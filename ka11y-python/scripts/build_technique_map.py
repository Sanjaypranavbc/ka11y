#!/usr/bin/env python3
"""
scripts/build_technique_map.py
==============================
Regenerate ``ka11y/data/wcag-technique-map.json`` from the reference workbook
``WCAG_Testing_Report_CodeCoverage.xlsx`` (sheet ``All Techniques``).

The JSON is the runtime source for ``ka11y.accessibility.technique_map``: it
tags every audit finding with the WCAG Situation(s) and Technique(s) that test
its success criterion. The spreadsheet is never read at runtime — rerun this
script and commit the JSON whenever the workbook changes::

    cd ka11y-python
    poetry run python scripts/build_technique_map.py \
        --xlsx ~/Downloads/WCAG_Testing_Report_CodeCoverage.xlsx

Columns read (by header name, so column order does not matter):

    SC | Situation | Technique | Technique Cover | Technique ID | Code Evidence

* ``Situation`` is forward-filled: the per-SC sheets only write it on the first
  row of each situation group (merged-cell style) and ``All Techniques`` is
  copied from them, so a blank cell means "same as the row above".
* Rows without a ``Technique ID`` are group headings ("Short text alternative
  techniques for Situation B:") or "(future technique)" placeholders. They are
  not techniques and are dropped; the count is reported in ``skipped``.

Rule correlation
----------------
Most checks do not record which technique they implement, so the finding →
technique match at runtime works from the spreadsheet's own ``Code Evidence``
column, which cites the source files behind each technique. This script
resolves those citations to the rule ids the engines emit:

* ``ka11y-node/src/custom-checks/<name>.check.js``  → rule id ``custom-<name>``
  (every custom check's RULE_ID follows that convention).
* ``ka11y-node/src/utils/rulesGuide.js:<line>``      → the axe-core rule whose
  guide entry contains that line (needs the checkout; ``--repo-root``).
* Any other ``ka11y-node`` path (axe/AccessLint configuration, the result
  mapper, the crawl service) → engine tag ``axe``.
* Any ``ka11y-python`` path → engine tag ``python`` (Python rule ids embed
  their SC, ``python_1_1_1_alt``, so SC + engine is enough).

Each technique entry therefore carries ``rules`` (exact rule ids), ``engines``
(coarser tags) and ``evidence`` (the cited files, for auditing the mapping).
"""

from __future__ import annotations

import argparse
import datetime as _dt
import json
import re
import sys
from collections import OrderedDict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

_HERE = Path(__file__).resolve().parent
_PKG_ROOT = _HERE.parent                   # ka11y-python/
_DEFAULT_OUT = _PKG_ROOT / "ka11y" / "data" / "wcag-technique-map.json"
#: WCAG 2.2 criteria the workbook does not cover; merged into the output (see file).
_DEFAULT_SUPPLEMENT = _PKG_ROOT / "ka11y" / "data" / "wcag-technique-supplement.json"
_DEFAULT_SHEET = "All Techniques"
_XLSX_NAME = "WCAG_Testing_Report_CodeCoverage.xlsx"
_XLSX_CANDIDATES = (
    Path.cwd() / _XLSX_NAME,
    _PKG_ROOT.parent / _XLSX_NAME,
    Path("/mnt/user-data/uploads") / _XLSX_NAME,
    Path.home() / "Downloads" / _XLSX_NAME,
)

_EVIDENCE_RE = re.compile(r"(ka11y-(?:node|python)/[\w./\-]+?\.(?:js|py|ts|yml|yaml|json))(?::(\d+))?")
_CUSTOM_CHECK_RE = re.compile(r"ka11y-node/src/custom-checks/([\w\-]+)\.check\.js$")
_RULES_GUIDE_PATH = "ka11y-node/src/utils/rulesGuide.js"
_GUIDE_KEY_RE = re.compile(r"^\s{2}'([a-z0-9\-]+)':\s*\{")
_COVER_CANON = {
    "implemented": "Implemented",
    "partially implemented": "Partially Implemented",
    "not implemented": "Not Implemented",
}


def _cell(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _sc_sort_key(sc: str) -> Tuple[int, ...]:
    try:
        return tuple(int(p) for p in sc.split("."))
    except ValueError:
        return (999,)


def _technique_name(raw: str, tid: str) -> str:
    """'Technique G94: Providing …' / 'ARIA6: Using …' → 'Providing …' / 'Using …'."""
    text = " ".join(raw.split())
    m = re.match(rf"^(?:Technique\s+)?{re.escape(tid)}\s*[:：]\s*", text, flags=re.IGNORECASE)
    return text[m.end():].strip() if m else text


def load_rules_guide_index(repo_root: Optional[Path]) -> List[Tuple[int, str]]:
    """[(start_line, axe_rule_id), …] for every entry in rulesGuide.js, or []."""
    if repo_root is None:
        return []
    path = repo_root / _RULES_GUIDE_PATH
    if not path.is_file():
        return []
    index: List[Tuple[int, str]] = []
    with path.open(encoding="utf-8") as fh:
        for n, line in enumerate(fh, start=1):
            m = _GUIDE_KEY_RE.match(line)
            if m:
                index.append((n, m.group(1)))
    return index


def resolve_guide_rule(index: List[Tuple[int, str]], line: int) -> Optional[str]:
    rule = None
    for start, rid in index:
        if start <= line:
            rule = rid
        else:
            break
    return rule


def correlate(evidence_text: str, guide_index: List[Tuple[int, str]]) -> Dict[str, List[str]]:
    """Code Evidence text → {"rules": [...], "engines": [...], "evidence": [...]}."""
    rules: "OrderedDict[str, None]" = OrderedDict()
    engines: "OrderedDict[str, None]" = OrderedDict()
    files: "OrderedDict[str, None]" = OrderedDict()
    for m in _EVIDENCE_RE.finditer(evidence_text or ""):
        path, line = m.group(1), m.group(2)
        files[path] = None
        cm = _CUSTOM_CHECK_RE.search(path)
        if cm:
            rules[f"custom-{cm.group(1)}"] = None
            engines["custom"] = None
        elif path == _RULES_GUIDE_PATH:
            engines["axe"] = None
            if line:
                rid = resolve_guide_rule(guide_index, int(line))
                if rid:
                    rules[rid] = None
        elif path.startswith("ka11y-node/"):
            engines["axe"] = None
        elif path.startswith("ka11y-python/"):
            engines["python"] = None
    return {"rules": list(rules), "engines": list(engines), "evidence": list(files)}


def read_rows(xlsx: Path, sheet: str) -> Tuple[List[str], List[List[Any]]]:
    try:
        import openpyxl  # noqa: WPS433 — optional dependency of the build step only
    except ImportError:  # pragma: no cover
        sys.exit("openpyxl is required: poetry run pip install openpyxl")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        sys.exit(f"sheet {sheet!r} not found in {xlsx} (have: {wb.sheetnames})")
    rows = list(wb[sheet].iter_rows(values_only=True))
    if not rows:
        sys.exit(f"sheet {sheet!r} is empty")
    header = [_cell(h) for h in rows[0]]
    return header, [list(r) for r in rows[1:]]


def build_map(
    header: List[str],
    rows: Iterable[List[Any]],
    *,
    guide_index: Optional[List[Tuple[int, str]]] = None,
    source: str = _XLSX_NAME,
    sheet: str = _DEFAULT_SHEET,
) -> Dict[str, Any]:
    """Pure transform: header + data rows → the JSON document (a dict)."""
    guide_index = guide_index or []

    def col(name: str, required: bool = True) -> Optional[int]:
        for i, h in enumerate(header):
            if h.lower() == name.lower():
                return i
        if required:
            sys.exit(f"column {name!r} missing from header {header}")
        return None

    c_sc, c_sit, c_tech = col("SC"), col("Situation"), col("Technique")
    c_cover, c_id = col("Technique Cover"), col("Technique ID")
    c_ev = col("Code Evidence", required=False)

    by_sc: Dict[str, List[Dict[str, Any]]] = {}
    seen_ids: Dict[Tuple[str, str], int] = {}
    skipped = 0
    last_situation = ""
    for row in rows:
        get = lambda i: _cell(row[i]) if i is not None and i < len(row) else ""  # noqa: E731
        sc = get(c_sc)
        situation = get(c_sit) or last_situation      # forward-fill merged cells
        if get(c_sit):
            last_situation = get(c_sit)
        tid = get(c_id)
        if not sc or not tid:
            if sc or get(c_tech):
                skipped += 1
            continue
        cover_raw = get(c_cover)
        cover = _COVER_CANON.get(cover_raw.lower(), cover_raw or "No Status")
        entry: Dict[str, Any] = {
            "id": tid,
            "name": _technique_name(get(c_tech), tid),
            "situation": situation,
            "cover": cover,
        }
        entry.update(correlate(get(c_ev), guide_index))
        key = (sc, tid)
        if key in seen_ids:
            # Same technique listed twice under one SC (different situations):
            # keep the first, remember the extra situation.
            first = by_sc[sc][seen_ids[key]]
            extra = first.setdefault("also_situations", [])
            if situation and situation != first["situation"] and situation not in extra:
                extra.append(situation)
            continue
        seen_ids[key] = len(by_sc.setdefault(sc, []))
        by_sc[sc].append(entry)

    ordered = OrderedDict((sc, by_sc[sc]) for sc in sorted(by_sc, key=_sc_sort_key))
    covers: Dict[str, int] = {}
    for entries in ordered.values():
        for e in entries:
            covers[e["cover"]] = covers.get(e["cover"], 0) + 1
    return OrderedDict(
        [
            ("schema_version", 1),
            ("source", source),
            ("sheet", sheet),
            ("generated_at", _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0).isoformat()),
            ("generator", "scripts/build_technique_map.py"),
            ("technique_count", sum(len(v) for v in ordered.values())),
            ("sc_count", len(ordered)),
            ("skipped", skipped),
            ("cover_counts", covers),
            ("by_sc", ordered),
        ]
    )


def merge_supplement(doc: Dict[str, Any], supplement_path: Optional[Path]) -> int:
    """Append the supplement's SCs/techniques to *doc* (in place). Techniques the
    workbook already lists for an SC are kept as they are. Returns the number
    of techniques added."""
    if supplement_path is None or not supplement_path.is_file():
        return 0
    supp = json.loads(supplement_path.read_text(encoding="utf-8"))
    added = 0
    by_sc: Dict[str, List[Dict[str, Any]]] = doc["by_sc"]
    for sc, entries in (supp.get("by_sc") or {}).items():
        have = {e["id"] for e in by_sc.get(sc, [])}
        for e in entries:
            if e["id"] in have:
                continue
            by_sc.setdefault(sc, []).append(dict(e, supplement=True))
            have.add(e["id"])
            added += 1
    doc["by_sc"] = OrderedDict((sc, by_sc[sc]) for sc in sorted(by_sc, key=_sc_sort_key))
    doc["technique_count"] = sum(len(v) for v in doc["by_sc"].values())
    doc["sc_count"] = len(doc["by_sc"])
    doc["supplement"] = supplement_path.name
    doc["supplement_count"] = added
    covers: Dict[str, int] = {}
    for entries in doc["by_sc"].values():
        for e in entries:
            covers[e["cover"]] = covers.get(e["cover"], 0) + 1
    doc["cover_counts"] = covers
    return added


def _find_xlsx(explicit: Optional[str]) -> Path:
    if explicit:
        p = Path(explicit).expanduser()
        if not p.is_file():
            sys.exit(f"workbook not found: {p}")
        return p
    for cand in _XLSX_CANDIDATES:
        if cand.is_file():
            return cand
    sys.exit(f"{_XLSX_NAME} not found; pass --xlsx PATH")


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--xlsx", help=f"path to {_XLSX_NAME} (default: cwd, repo root, ~/Downloads)")
    ap.add_argument("--sheet", default=_DEFAULT_SHEET)
    ap.add_argument("--out", default=str(_DEFAULT_OUT), help="output JSON path")
    ap.add_argument("--supplement", default=str(_DEFAULT_SUPPLEMENT),
                    help="WCAG 2.2 supplement JSON merged into the output ('' to skip)")
    ap.add_argument(
        "--repo-root",
        default=str(_PKG_ROOT.parent),
        help="monorepo root holding ka11y-node/ (resolves rulesGuide.js line citations)",
    )
    args = ap.parse_args(argv)

    xlsx = _find_xlsx(args.xlsx)
    header, rows = read_rows(xlsx, args.sheet)
    guide_index = load_rules_guide_index(Path(args.repo_root).expanduser())
    doc = build_map(header, rows, guide_index=guide_index, source=xlsx.name, sheet=args.sheet)
    added = merge_supplement(doc, Path(args.supplement).expanduser() if args.supplement else None)

    out = Path(args.out).expanduser()
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(doc, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    with_rules = sum(1 for es in doc["by_sc"].values() for e in es if e["rules"])
    print(
        f"{out}: {doc['technique_count']} techniques across {doc['sc_count']} SCs "
        f"({with_rules} with exact rule correlation, {doc['skipped']} placeholder rows skipped; "
        f"rulesGuide index: {len(guide_index)} rules; supplement: {added} techniques)"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
